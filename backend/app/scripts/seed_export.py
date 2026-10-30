"""Write the seeded data to an .xlsx file, one sheet per topic."""

from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.modules.accounts.models import Account, Transaction
from app.modules.classrooms.models import ChildClassMembership, ClassAccessCode, SchoolClass
from app.modules.fundraisers.models import Fundraiser, FundraiserObligation, FundraiserPayment
from app.modules.users.models import Child, User

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E6A4C")


def _cell_value(value: Any) -> Any:
    if isinstance(value, UUID):
        return str(value)
    if hasattr(value, "value") and isinstance(value.value, str):  # enums
        return value.value
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)  # Excel has no time zones
    return value


def _add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for row in rows:
        sheet.append([_cell_value(value) for value in row])

    for column in sheet.iter_cols(min_row=2):
        for cell in column:
            if isinstance(cell.value, Decimal | float):
                cell.number_format = MONEY_FORMAT
            elif isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT

    for index, column in enumerate(sheet.columns, start=1):
        longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 60)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_xlsx(db: Session, path: Path, password: str, logins: list[tuple[str, str]]) -> None:
    users = {user.id: user for user in db.scalars(select(User))}
    classes = {school_class.id: school_class for school_class in db.scalars(select(SchoolClass))}
    accounts = {account.id: account for account in db.scalars(select(Account))}
    fundraisers = {fundraiser.id: fundraiser for fundraiser in db.scalars(select(Fundraiser))}

    def owner(account_id: UUID | None) -> str:
        if account_id is None:
            return "(outside the system)"
        account = accounts[account_id]
        if account.user_id is not None:
            return users[account.user_id].email
        return f"fundraiser: {fundraisers[account.fundraiser_id].name}"

    workbook = Workbook()
    workbook.remove(workbook.active)

    # Logins
    by_email = {user.email: user for user in users.values()}
    _add_sheet(
        workbook,
        "Logins",
        ["Email", "Password", "What it is for", "Name", "Role", "Status"],
        [
            [
                email,
                password,
                purpose,
                f"{by_email[email].first_name} {by_email[email].last_name}",
                by_email[email].role,
                by_email[email].status,
            ]
            for email, purpose in logins
        ],
    )

    # Wallets
    _add_sheet(
        workbook,
        "Wallets",
        ["Email", "Name", "IBAN", "Balance (PLN)"],
        [
            [
                users[account.user_id].email,
                f"{users[account.user_id].first_name} {users[account.user_id].last_name}",
                account.iban,
                account.balance,
            ]
            for account in sorted(
                (a for a in accounts.values() if a.user_id is not None),
                key=lambda a: users[a.user_id].email,
            )
        ],
    )

    # Classes
    codes = {code.class_id: code.code for code in db.scalars(select(ClassAccessCode))}
    member_counts = dict(
        db.execute(
            select(ChildClassMembership.class_id, func.count())
            .where(ChildClassMembership.status == "ACTIVE")
            .group_by(ChildClassMembership.class_id)
        ).all()
    )
    _add_sheet(
        workbook,
        "Classes",
        ["Class", "School", "Treasurer", "Access code", "Children (active)"],
        [
            [
                c.name,
                c.school_name,
                users[c.treasurer_id].email,
                codes.get(c.id),
                member_counts.get(c.id, 0),
            ]
            for c in sorted(classes.values(), key=lambda c: c.name)
        ],
    )

    # Children
    children = {child.id: child for child in db.scalars(select(Child))}
    _add_sheet(
        workbook,
        "Children",
        ["Child", "Birth date", "Parent", "Class", "Membership", "Joined (UTC)"],
        [
            [
                f"{children[m.child_id].first_name} {children[m.child_id].last_name}",
                children[m.child_id].birth_date,
                users[children[m.child_id].parent_id].email,
                classes[m.class_id].name,
                m.status,
                m.joined_at,
            ]
            for m in sorted(
                db.scalars(select(ChildClassMembership)),
                key=lambda m: (
                    classes[m.class_id].name,
                    users[children[m.child_id].parent_id].email,
                ),
            )
        ],
    )

    # Fundraisers
    obligations = dict(
        db.execute(
            select(FundraiserObligation.fundraiser_id, func.count()).group_by(
                FundraiserObligation.fundraiser_id
            )
        ).all()
    )
    paid = dict(
        db.execute(
            select(FundraiserPayment.fundraiser_id, func.count())
            .where(FundraiserPayment.refunded_at.is_(None))
            .group_by(FundraiserPayment.fundraiser_id)
        ).all()
    )
    fundraiser_balance = {a.fundraiser_id: a.balance for a in accounts.values() if a.fundraiser_id}
    _add_sheet(
        workbook,
        "Fundraisers",
        [
            "Fundraiser",
            "Class",
            "Status",
            "Per child (PLN)",
            "Paid / owed",
            "Account balance (PLN)",
            "Start",
            "End",
        ],
        [
            [
                f.name,
                classes[f.class_id].name,
                f.status,
                f.required_amount,
                f"{paid.get(f.id, 0)} / {obligations.get(f.id, 0)}",
                fundraiser_balance.get(f.id),
                f.start_date,
                f.end_date,
            ]
            for f in sorted(
                fundraisers.values(), key=lambda f: (classes[f.class_id].name, f.start_date)
            )
        ],
    )

    # Payments
    _add_sheet(
        workbook,
        "Payments",
        ["Fundraiser", "Child", "Paid by", "Amount (PLN)", "Paid at (UTC)", "Refunded at (UTC)"],
        [
            [
                fundraisers[p.fundraiser_id].name,
                f"{children[p.child_id].first_name} {children[p.child_id].last_name}",
                users[p.payer_id].email,
                p.amount,
                p.paid_at,
                p.refunded_at,
            ]
            for p in sorted(
                db.scalars(select(FundraiserPayment)),
                key=lambda p: (fundraisers[p.fundraiser_id].name, p.paid_at),
            )
        ],
    )

    # Transactions (the ledger)
    _add_sheet(
        workbook,
        "Transactions",
        ["Created at (UTC)", "Type", "From", "To", "Amount (PLN)", "Description"],
        [
            [
                t.created_at,
                t.type,
                owner(t.from_account_id),
                owner(t.to_account_id),
                t.amount,
                t.description,
            ]
            for t in db.scalars(select(Transaction).order_by(Transaction.created_at))
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
